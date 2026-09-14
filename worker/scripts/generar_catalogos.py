"""Genera los catálogos semilla de CNAE y municipios desde los ficheros
oficiales del INE.

Escribe tres CSV en `supabase/seed/` (los que planea el README de esa
carpeta):

    cnae.csv                  estructura de CNAE-2009 y CNAE-2025
    cnae_correspondencias.csv correspondencias entre ambas versiones
    municipios.csv            municipios con código INE, provincia y CCAA

No toca la base de datos: para cargarlos, `scripts/cargar_catalogos.py`.
Requiere haber aplicado antes la migración 202609140001, que añade la
clave primaria compuesta de `cnae` y la tabla de correspondencias.

Uso (con el entorno virtual del worker activado):
    python scripts\\generar_catalogos.py
    python scripts\\generar_catalogos.py --sin-descarga   # reutiliza la caché

Necesita `openpyxl`, que está en las dependencias `dev` del worker
(`pip install -e ".[dev]"`): es un script de mantenimiento que se ejecuta
muy de vez en cuando, no forma parte del runtime que se despliega.

Fuentes (verificadas el 2026-09-14):
    ine.es/daco/daco42/clasificaciones/cnae25/Estructura_CNAE2025.xlsx
    ine.es/daco/daco42/clasificaciones/cnae25/Correspondencia_CNAE09_CNAE25.xlsx
    ine.es/daco/daco42/clasificaciones/cnae25/Correspondencia_CNAE25_CNAE09.xlsx
    ine.es/daco/daco42/codmun/26codmun.xlsx
    ine.es/daco/daco42/codmun/cod_ccaa_provincia.htm
"""

from __future__ import annotations

import argparse
import csv
import html
import re
import sys
import unicodedata
from collections.abc import Iterable, Iterator
from pathlib import Path

import httpx
import openpyxl

RAIZ = Path(__file__).resolve().parents[2]  # .../radar-b2b
SALIDA = RAIZ / "supabase" / "seed"
CACHE = SALIDA / "_ine"

BASE_CNAE = "https://www.ine.es/daco/daco42/clasificaciones/cnae25"
BASE_MUN = "https://www.ine.es/daco/daco42/codmun"

FICHEROS: dict[str, str] = {
    "Estructura_CNAE2025.xlsx": f"{BASE_CNAE}/Estructura_CNAE2025.xlsx",
    "Correspondencia_CNAE09_CNAE25.xlsx": f"{BASE_CNAE}/Correspondencia_CNAE09_CNAE25.xlsx",
    "Correspondencia_CNAE25_CNAE09.xlsx": f"{BASE_CNAE}/Correspondencia_CNAE25_CNAE09.xlsx",
    "26codmun.xlsx": f"{BASE_MUN}/26codmun.xlsx",
    "cod_ccaa_provincia.htm": f"{BASE_MUN}/cod_ccaa_provincia.htm",
}

V09 = "CNAE-2009"
V25 = "CNAE-2025"

AGENTE = "Mozilla/5.0 (compatible; RadarB2B/1.0; +uso interno Solventa IA)"


# ---------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------
def normalizar_texto(t: str) -> str:
    """Replica en Python la función `normalizar_texto` de Postgres.

    `unaccent` convierte la ñ en n, así que aquí también (coherente con el
    doc 05 §2.3). Verificado contra la BD: de los 8.132 municipios, ninguno
    conserva la ñ en `nombre_norm`.
    """
    sin_tildes = unicodedata.normalize("NFKD", str(t))
    sin_tildes = "".join(c for c in sin_tildes if not unicodedata.combining(c))
    limpio = re.sub(r"[^a-z0-9ñ ]", " ", sin_tildes.lower())
    return re.sub(r"\s+", " ", limpio).strip()


def limpiar_codigo(valor: object) -> str:
    """'41.01' -> '4101'; 1 -> '01' (división leída por openpyxl como int)."""
    if isinstance(valor, int):
        return f"{valor:02d}"
    texto = str(valor).strip().replace(".", "")
    if texto.isdigit() and len(texto) == 1:
        return texto.zfill(2)
    return texto


def nivel_de(codigo: str) -> int:
    """1 sección (letra), 2 división, 3 grupo, 4 clase."""
    if not codigo.isdigit():
        return 1
    return {2: 2, 3: 3, 4: 4}[len(codigo)]


def descargar(sin_descarga: bool) -> None:
    CACHE.mkdir(parents=True, exist_ok=True)
    with httpx.Client(timeout=60.0, headers={"User-Agent": AGENTE}, follow_redirects=True) as cli:
        for nombre, url in FICHEROS.items():
            ruta = CACHE / nombre
            if sin_descarga and ruta.exists():
                print(f"  (caché) {nombre}")
                continue
            respuesta = cli.get(url)
            respuesta.raise_for_status()
            ruta.write_bytes(respuesta.content)
            print(f"  descargado {nombre} ({len(respuesta.content):,} bytes)")


def escribir_csv(nombre: str, cabecera: list[str], filas: Iterable[tuple]) -> None:
    ruta = SALIDA / nombre
    with ruta.open("w", encoding="utf-8", newline="") as f:
        escritor = csv.writer(f)
        escritor.writerow(cabecera)
        escritor.writerows(filas)
    print(f"  escrito {nombre} ({ruta.stat().st_size:,} bytes)")


def filas_hoja(nombre_fichero: str, hoja: str) -> Iterator[tuple]:
    wb = openpyxl.load_workbook(CACHE / nombre_fichero, read_only=True, data_only=True)
    try:
        yield from wb[hoja].iter_rows(values_only=True)
    finally:
        wb.close()


# ---------------------------------------------------------------------
# CNAE: estructura
# ---------------------------------------------------------------------
def estructura_desde(pares: Iterable[tuple[str, str]]) -> list[tuple[str, str, int, str | None]]:
    """Convierte (código, título) en orden jerárquico a filas con padre.

    Los ficheros del INE vienen ordenados: una sección y a continuación sus
    divisiones, grupos y clases. Se aprovecha ese orden para saber a qué
    sección pertenece cada división, que es el único salto que no se puede
    deducir del propio código.
    """
    salida: list[tuple[str, str, int, str | None]] = []
    vistos: set[str] = set()
    seccion_actual: str | None = None

    for codigo, titulo in pares:
        if not codigo or codigo in vistos:
            continue
        vistos.add(codigo)
        nivel = nivel_de(codigo)
        if nivel == 1:
            seccion_actual = codigo
            padre = None
        elif nivel == 2:
            padre = seccion_actual
        else:
            padre = codigo[: nivel - 1]
        salida.append((codigo, titulo, nivel, padre))

    salida.sort(key=lambda f: (f[2], f[0]))  # por nivel: los padres primero
    return salida


def leer_estructura_2025() -> list[tuple[str, str, int, str | None]]:
    pares: list[tuple[str, str]] = []
    for fila in filas_hoja("Estructura_CNAE2025.xlsx", "Estructura_CNAE2025"):
        if not fila or fila[0] is None or fila[1] is None:
            continue
        codigo = limpiar_codigo(fila[0])
        if codigo.startswith("CÓDIGO"):
            continue
        pares.append((codigo, str(fila[1]).strip()))
    return estructura_desde(pares)


def leer_estructura_2009() -> list[tuple[str, str, int, str | None]]:
    """CNAE-2009 se extrae del fichero de correspondencias.

    Verificado el 2026-09-14: contiene las 1.010 rúbricas completas de
    CNAE-2009 (21 secciones, 88 divisiones, 272 grupos, 629 clases), así que
    no hace falta el `.xls` de estructura de 2009 — y con él nos ahorramos
    la dependencia `xlrd`.
    """
    pares: list[tuple[str, str]] = []
    for fila in filas_hoja("Correspondencia_CNAE09_CNAE25.xlsx", "Correspond_CNAE09_CNAE25"):
        if not fila or fila[0] is None or fila[1] is None:
            continue
        codigo = limpiar_codigo(fila[0])
        if codigo.startswith("CNAE"):
            continue
        pares.append((codigo, str(fila[1]).strip()))
    return estructura_desde(pares)


# ---------------------------------------------------------------------
# CNAE: correspondencias
# ---------------------------------------------------------------------
def leer_correspondencias() -> list[tuple[str, str, str, str, int]]:
    ficheros = [
        ("Correspondencia_CNAE09_CNAE25.xlsx", "Correspond_CNAE09_CNAE25", V09, V25),
        ("Correspondencia_CNAE25_CNAE09.xlsx", "Correspondencias_CNAE25_CNAE09", V25, V09),
    ]
    vistos: set[tuple[str, str, str, str]] = set()
    salida: list[tuple[str, str, str, str, int]] = []

    for nombre, hoja, v_origen, v_destino in ficheros:
        for fila in filas_hoja(nombre, hoja):
            if not fila or fila[0] is None or fila[2] is None or not isinstance(fila[4], int):
                continue
            origen = limpiar_codigo(fila[0])
            destino = limpiar_codigo(fila[2])
            clave = (origen, v_origen, destino, v_destino)
            if clave in vistos:
                continue
            vistos.add(clave)
            salida.append((origen, v_origen, destino, v_destino, int(fila[4])))
    return salida


# ---------------------------------------------------------------------
# Municipios
# ---------------------------------------------------------------------
def leer_ccaa() -> dict[str, str]:
    bruto = (CACHE / "cod_ccaa_provincia.htm").read_bytes()
    texto = bruto.decode("iso-8859-1", errors="replace")
    mapa: dict[str, str] = {}
    for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", texto, re.DOTALL | re.IGNORECASE):
        celdas = [
            html.unescape(re.sub(r"<[^>]+>", "", c)).strip()
            for c in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", tr, re.DOTALL | re.IGNORECASE)
        ]
        if len(celdas) >= 4 and re.fullmatch(r"\d{2}", celdas[2]):
            mapa[celdas[2]] = celdas[1]
    if len(mapa) != 52:
        raise SystemExit(
            f"CCAA: esperaba 52 provincias, encontré {len(mapa)}. "
            "El INE ha cambiado el HTML: revisa cod_ccaa_provincia.htm."
        )
    return mapa


def leer_municipios() -> list[tuple[str, str, str, str, str]]:
    ccaa_por_provincia = leer_ccaa()
    wb = openpyxl.load_workbook(CACHE / "26codmun.xlsx", read_only=True, data_only=True)
    salida: list[tuple[str, str, str, str, str]] = []

    try:
        for hoja in wb.sheetnames:
            filas = list(wb[hoja].iter_rows(values_only=True))
            nombre_provincia = hoja
            for fila in filas[:3]:
                primera = str(fila[0]).strip() if fila and fila[0] is not None else ""
                if primera and primera != "CPRO" and not primera.startswith("Relación"):
                    nombre_provincia = primera
                    break
            for fila in filas:
                if not fila or fila[0] is None or fila[1] is None:
                    continue
                cpro, cmun = str(fila[0]).strip(), str(fila[1]).strip()
                if not re.fullmatch(r"\d{2}", cpro) or not re.fullmatch(r"\d{3}", cmun):
                    continue
                salida.append(
                    (cpro + cmun, str(fila[3]).strip(), nombre_provincia, cpro, ccaa_por_provincia[cpro])
                )
    finally:
        wb.close()

    # El lookup `buscar_municipio_ine` depende de que (provincia, nombre
    # normalizado) sea único. Verificado el 2026-09-14: 0 colisiones en toda
    # España. Si el INE renombra un municipio y rompe eso, mejor fallar aquí
    # que cargar datos que hagan ambiguo el lookup.
    claves = [(m[3], normalizar_texto(m[1])) for m in salida]
    if len(claves) != len(set(claves)):
        repetidas = {c for c in claves if claves.count(c) > 1}
        raise SystemExit(f"Colisión de nombre normalizado dentro de una provincia: {repetidas}")
    return salida


# ---------------------------------------------------------------------
def main() -> int:
    parser = argparse.ArgumentParser(description="Genera los catálogos CNAE y municipios desde el INE.")
    parser.add_argument("--sin-descarga", action="store_true", help="reutiliza los ficheros ya descargados")
    args = parser.parse_args()

    SALIDA.mkdir(parents=True, exist_ok=True)

    print("Descargando ficheros del INE...")
    descargar(args.sin_descarga)

    print("Leyendo estructura CNAE...")
    e09 = leer_estructura_2009()
    e25 = leer_estructura_2025()
    print(f"  {V09}: {len(e09)} rúbricas | {V25}: {len(e25)} rúbricas")

    print("Leyendo correspondencias...")
    correspondencias = leer_correspondencias()
    print(f"  {len(correspondencias)} correspondencias (las dos direcciones)")

    print("Leyendo municipios...")
    municipios = leer_municipios()
    print(f"  {len(municipios)} municipios")

    print("Escribiendo CSV en supabase/seed/...")
    escribir_csv(
        "cnae.csv",
        ["codigo", "version", "descripcion", "nivel", "codigo_padre"],
        [(c, V09, d, n, p or "") for c, d, n, p in e09] + [(c, V25, d, n, p or "") for c, d, n, p in e25],
    )
    escribir_csv(
        "cnae_correspondencias.csv",
        ["codigo_origen", "version_origen", "codigo_destino", "version_destino", "nivel"],
        correspondencias,
    )
    escribir_csv(
        "municipios.csv",
        ["codigo_ine", "nombre", "provincia", "cod_provincia", "ccaa"],
        municipios,
    )

    print("\nListo. Cargar con: python scripts\\cargar_catalogos.py")
    return 0


if __name__ == "__main__":
    sys.exit(main())
